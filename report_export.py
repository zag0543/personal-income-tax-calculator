"""
报告导出模块 — 生成 HTML / Word / PDF 格式的个税报告
"""

import json
from datetime import datetime

from scripts.deduction_checkup import DEDUCTION_POLICIES

# ── HTML 模板 ──

HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
  body {{ font-family: "Microsoft YaHei", "PingFang SC", sans-serif; max-width: 800px; margin: 0 auto; padding: 20px; color: #333; }}
  h1 {{ text-align: center; color: #1a1a2e; border-bottom: 2px solid #e94560; padding-bottom: 10px; }}
  h2 {{ color: #1a1a2e; margin-top: 24px; border-left: 4px solid #e94560; padding-left: 10px; }}
  .meta {{ text-align: center; color: #888; font-size: 14px; }}
  table {{ width: 100%; border-collapse: collapse; margin: 12px 0; }}
  th {{ background: #1a1a2e; color: #fff; padding: 8px 10px; text-align: left; }}
  td {{ padding: 8px 10px; border-bottom: 1px solid #eee; }}
  tr:hover td {{ background: #f5f5f5; }}
  .highlight {{ color: #e94560; font-weight: bold; }}
  .success {{ color: #2ecc71; font-weight: bold; }}
  .summary {{ background: #f8f9fa; border: 1px solid #ddd; border-radius: 8px; padding: 16px; margin: 12px 0; }}
  .summary-item {{ display: inline-block; margin: 8px 16px; text-align: center; }}
  .summary-item .label {{ font-size: 12px; color: #888; }}
  .summary-item .value {{ font-size: 20px; font-weight: bold; color: #1a1a2e; }}
  .footer {{ text-align: center; color: #aaa; font-size: 12px; margin-top: 32px; border-top: 1px solid #eee; padding-top: 12px; }}
</style>
</head>
<body>
<h1>{title}</h1>
<p class="meta">生成日期：{date} ｜ 数据仅供页面参考，实际以税务机关核定为准</p>
{content}
<div class="footer">
  <p>由「个人所得税计算器」生成 · 数据仅在浏览器本地计算</p>
</div>
</body>
</html>"""


def fmt(val):
    """格式化金额"""
    return f"¥{val:,.2f}"


def build_table(headers, rows):
    """构建 HTML 表格"""
    thead = "<tr>" + "".join(f"<th>{h}</th>" for h in headers) + "</tr>"
    tbody = ""
    for row in rows:
        tbody += "<tr>" + "".join(f"<td>{c}</td>" for c in row) + "</tr>"
    return f"<table>{thead}{tbody}</table>"


def gen_detailed_report(result, monthly_details=None):
    """生成详细计算报告 HTML"""
    now = datetime.now().strftime("%Y年%m月%d日 %H:%M")

    # 核心指标
    summary_html = f"""
    <div class="summary">
      <div class="summary-item"><div class="label">应纳税所得额</div><div class="value">{fmt(result['taxable_income'])}</div></div>
      <div class="summary-item"><div class="label">综合所得税额</div><div class="value">{fmt(result['income_tax'])}</div></div>
      <div class="summary-item"><div class="label">年终奖税额</div><div class="value">{fmt(result['bonus_tax'])}</div></div>
      <div class="summary-item"><div class="label">合计应纳税额</div><div class="value highlight">{fmt(result['total_tax'])}</div></div>
    </div>
    """

    # 收入与扣除明细
    other_ded = result['total_deductions'] - 60000
    detail_rows = [
        ("年度总收入", fmt(result['annual_income'])),
        ("基本减除费用", fmt(60000)),
        ("专项扣除及其他扣除合计", fmt(max(other_ded, 0))),
        ("扣除项目总计", fmt(result['total_deductions'])),
        ("应纳税所得额", fmt(result['taxable_income'])),
    ]
    detail_table = build_table(["项目", "金额"], detail_rows)

    # 税额明细
    tax_rows = [
        ("综合所得税额", fmt(result['income_tax'])),
        ("年终奖税额", fmt(result['bonus_tax'])),
        ("合计应纳税额", fmt(result['total_tax'])),
    ]
    if result.get('unit_withheld', 0) > 0:
        tax_rows.append(("单位已代扣", fmt(result['unit_withheld'])))
        tax_rows.append(("汇算结果", f"{result['result']}：{fmt(result['amount'])}"))
    tax_table = build_table(["项目", "金额"], tax_rows)

    # 月度明细
    monthly_html = ""
    if monthly_details:
        m_rows = []
        for m in monthly_details:
            m_rows.append((str(m['month']), fmt(m['monthly_income']), fmt(m['cumulative_income']), fmt(m['cumulative_deduction']), fmt(m['cumulative_taxable']), fmt(m['cumulative_tax']), fmt(m['monthly_tax'])))
        monthly_html = "<h2>月度预扣预缴明细</h2>" + build_table(
            ["月份", "本月收入", "累计收入", "累计扣除", "累计应纳税所得额", "累计应纳税额", "本月预扣"],
            m_rows,
        )

    content = f"""
    <h2>计算结果摘要</h2>
    {summary_html}
    <h2>收入与扣除明细</h2>
    {detail_table}
    <h2>税额明细</h2>
    {tax_table}
    {monthly_html}
    """

    return HTML_TEMPLATE.format(
        title="个人所得税年度汇算清缴报告",
        date=now,
        content=content,
    )


def gen_checkup_report(collector):
    """生成扣除项体检报告 HTML"""
    now = datetime.now().strftime("%Y年%m月%d日 %H:%M")
    diagnosis = collector.diagnosis
    data = collector.data
    tax_rate = data.get('estimated_tax_rate', 0.20)

    # 收入描述
    income = data.get('annual_income', 0)
    if isinstance(income, (list, tuple)):
        if income[1] == float('inf'):
            income_desc = f"{income[0]//10000}万元以上"
        else:
            income_desc = f"{income[0]//10000}-{income[1]//10000}万元"
    else:
        income_desc = "未填写"

    # 总览
    summary_html = f"""
    <div class="summary">
      <div class="summary-item"><div class="label">年收入</div><div class="value">{income_desc}</div></div>
      <div class="summary-item"><div class="label">适用税率（预估）</div><div class="value">{tax_rate*100:.0f}%</div></div>
      <div class="summary-item"><div class="label">可能遗漏</div><div class="value highlight">{fmt(diagnosis['total_missed'])}</div></div>
      <div class="summary-item"><div class="label">预计可节税</div><div class="value success">{fmt(diagnosis['potential_savings'])}</div></div>
    </div>
    """

    # 遗漏项
    missed_html = ""
    if diagnosis['possible_missed']:
        rows = []
        for item in diagnosis['possible_missed']:
            ptype = item['type']
            policy = DEDUCTION_POLICIES.get(ptype, {})
            name = policy.get('name', ptype)
            rows.append((name, fmt(item['amount']), fmt(item['savings']), item['action']))
        missed_html = "<h2>⚠️ 可能有遗漏的扣除项</h2>" + build_table(
            ["扣除项目", "预计扣除/年", "节税效果/年", "操作指引"],
            rows,
        )

    # 不符合项
    na_html = ""
    if diagnosis['not_applicable']:
        items = "".join(
            f"<li><strong>{DEDUCTION_POLICIES.get(item['type'], {}).get('name', item['type'])}</strong>：{item['reason']} → {item['suggestion']}</li>"
            for item in diagnosis['not_applicable']
        )
        na_html = f"<h2>❓ 暂不符合条件</h2><ul>{items}</ul>"

    # 行动清单
    action_html = ""
    if diagnosis['possible_missed']:
        items = "".join(
            f"<li><strong>{DEDUCTION_POLICIES.get(item['type'], {}).get('name', item['type'])}</strong>：{item['action']}</li>"
            for item in diagnosis['possible_missed']
        )
        action_html = f"<h2>📋 补申报行动清单</h2><ol>{items}</ol>"

    content = f"""
    <h2>健康总览</h2>
    {summary_html}
    {missed_html}
    {na_html}
    {action_html}
    <p>💡 温馨提示：扣除项填报可通过 <strong>个人所得税APP</strong> 完成，每年12月可确认次年扣除信息。</p>
    """

    return HTML_TEMPLATE.format(
        title="个税扣除项体检报告",
        date=now,
        content=content,
    )


def gen_joint_report(a_data, b_data, a_opt, b_opt, a_base, b_base,
                     base_total, opt_total, savings,
                     a_portion, b_portion, a_income, b_income):
    """生成夫妻共同申报报告 HTML"""
    now = datetime.now().strftime("%Y年%m月%d日 %H:%M")

    # 总体对比
    summary_html = f"""
    <div class="summary">
      <div class="summary-item"><div class="label">独立申报总税额</div><div class="value">{fmt(base_total)}</div></div>
      <div class="summary-item"><div class="label">优化分配总税额</div><div class="value success">{fmt(opt_total)}</div></div>
      <div class="summary-item"><div class="label">节省税款</div><div class="value highlight">{fmt(savings)}</div></div>
      <div class="summary-item"><div class="label">可分配扣除总额</div><div class="value">{fmt(a_portion + b_portion)}</div></div>
    </div>
    """

    # 双方对比
    comp_rows = [
        ("配偶 A", f"¥{a_income:,}/月", fmt(a_data['total_tax']), fmt(a_opt['total_tax']), fmt(a_portion)),
        ("配偶 B", f"¥{b_income:,}/月", fmt(b_data['total_tax']), fmt(b_opt['total_tax']), fmt(b_portion)),
    ]
    comp_table = build_table(
        ["成员", "月收入", "独立申报", "优化后", "获得扣除额"],
        comp_rows,
    )

    # 配偶明细
    a_detail = build_table(
        ["项目", "金额"],
        [
            ("月薪", f"¥{a_income:,}"),
            ("应纳税所得额", fmt(a_opt['taxable_income'])),
            ("应纳税额", fmt(a_opt['total_tax'])),
            ("分配扣除额", fmt(a_portion)),
        ]
    )
    b_detail = build_table(
        ["项目", "金额"],
        [
            ("月薪", f"¥{b_income:,}"),
            ("应纳税所得额", fmt(b_opt['taxable_income'])),
            ("应纳税额", fmt(b_opt['total_tax'])),
            ("分配扣除额", fmt(b_portion)),
        ]
    )

    content = f"""
    <h2>总体对比</h2>
    {summary_html}
    <h2>双方独立申报 vs 优化分配</h2>
    {comp_table}
    <h2>配偶 A 明细</h2>
    {a_detail}
    <h2>配偶 B 明细</h2>
    {b_detail}
    """

    return HTML_TEMPLATE.format(
        title="夫妻共同申报测算报告",
        date=now,
        content=content,
    )


def try_gen_docx(html_content, output_path):
    """尝试使用 python-docx 生成 Word 文件（可选依赖）"""
    try:
        from docx import Document
        from docx.shared import Inches
        import re
        from bs4 import BeautifulSoup

        doc = Document()
        soup = BeautifulSoup(html_content, 'html.parser')

        for tag in soup.find_all(['h1', 'h2', 'p', 'table']):
            if tag.name == 'h1':
                doc.add_heading(tag.get_text(), level=1)
            elif tag.name == 'h2':
                doc.add_heading(tag.get_text(), level=2)
            elif tag.name == 'p':
                doc.add_paragraph(tag.get_text())
            elif tag.name == 'table':
                rows = tag.find_all('tr')
                if rows:
                    cols = len(rows[0].find_all(['th', 'td']))
                    table = doc.add_table(rows=len(rows), cols=cols)
                    table.style = 'Light Grid Accent 1'
                    for i, row in enumerate(rows):
                        cells = row.find_all(['th', 'td'])
                        for j, cell in enumerate(cells):
                            table.rows[i].cells[j].text = cell.get_text()

        doc.save(output_path)
        return True
    except ImportError:
        return False
    except Exception:
        return False


def try_gen_pdf(html_content, output_path):
    """尝试使用 fpdf2 生成 PDF 文件（可选依赖）"""
    try:
        from fpdf import FPDF
        import re
        from bs4 import BeautifulSoup

        pdf = FPDF()
        pdf.add_page()
        pdf.add_font('zh', '', 'C:/Windows/Fonts/msyh.ttc', uni=True)
        pdf.set_font('zh', '', 12)

        soup = BeautifulSoup(html_content, 'html.parser')
        for tag in soup.find_all(['h1', 'h2', 'p', 'li']):
            text = tag.get_text(strip=True)
            if tag.name == 'h1':
                pdf.set_font('zh', '', 18)
                pdf.cell(0, 12, text, ln=True, align='C')
                pdf.set_font('zh', '', 12)
            elif tag.name == 'h2':
                pdf.set_font('zh', '', 14)
                pdf.cell(0, 10, text, ln=True)
                pdf.set_font('zh', '', 12)
            else:
                pdf.multi_cell(0, 8, text)

        pdf.output(output_path)
        return True
    except ImportError:
        return False
    except Exception:
        return False
